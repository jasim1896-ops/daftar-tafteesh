import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.formatting.rule import FormulaRule
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment

import sys
CLEAN = "--clean" in sys.argv
import os
HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "template_clean.xlsx") if CLEAN else os.path.join(HERE, "سجل_التفتيش_البيئي.xlsx")
N = 1500                      # rows available in the log
R0 = 5                        # first data row
R1 = R0 + N - 1               # last data row

FONT = "Arial"
GREEN = "1E5B3C"
GREEN2 = "2E7D52"
LIGHT = "E9F3EC"
SAND = "F7F1E3"
INPUT = "FFF9DB"
GREY = "ECEFF1"
WHITE = "FFFFFF"
thin = Side(style="thin", color="C9D3CC")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

def f(size=11, bold=False, color="222222"):
    return Font(name=FONT, size=size, bold=bold, color=color)

def fill(c):
    return PatternFill("solid", fgColor=c)

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center", wrap_text=True)

def style(ws, rng, font=None, fl=None, al=None, border=True, fmt=None):
    for row in ws[rng]:
        for c in row:
            if font: c.font = font
            if fl: c.fill = fl
            if al: c.alignment = al
            if border: c.border = BORDER
            if fmt: c.number_format = fmt

def banner(ws, rng, text, size=16):
    ws.merge_cells(rng)
    c = ws[rng.split(":")[0]]
    c.value = text
    style(ws, rng, f(size, True, WHITE), fill(GREEN), CENTER, border=False)

def section(ws, rng, text):
    ws.merge_cells(rng)
    c = ws[rng.split(":")[0]]
    c.value = text
    style(ws, rng, f(12, True, GREEN), fill(LIGHT), RIGHT, border=False)

wb = Workbook()

# ------------------------------------------------------------------ sheets
ws_help = wb.active; ws_help.title = "طريقة الاستخدام"
ws_log = wb.create_sheet("سجل الزيارات")
ws_w = wb.create_sheet("التقرير الأسبوعي")
ws_m = wb.create_sheet("التقرير الشهري")
ws_p = wb.create_sheet("التقرير الدوري")
ws_s = wb.create_sheet("الإعدادات")
for ws in wb.worksheets:
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False

S = "'الإعدادات'"
LOG = "'سجل الزيارات'"

# ------------------------------------------------------------------ settings
ws = ws_s
banner(ws, "A1:J1", "الإعدادات والقوائم")
labels = [("اسم المفتش", "جاسم الحمادي"),
          ("الجهة", "هيئة البيئة"),
          ("القطاع", "منطقة الظفرة"),
          ("فولدر الصور (بجانب الملف)", "الصور")]
for i, (k, v) in enumerate(labels):
    r = 3 + i
    ws.cell(r, 1, k); ws.cell(r, 2, v)
    style(ws, f"A{r}:A{r}", f(11, True), fill(LIGHT), RIGHT)
    style(ws, f"B{r}:B{r}", f(11, False, "1A3FA8"), fill(INPUT), RIGHT)
    ws.merge_cells(f"B{r}:C{r}")

section(ws, "A9:J9", "المواقع حسب المنطقة — لإضافة موقع جديد اكتبه في أول خانة فاضية تحت منطقته، ولإضافة منطقة جديدة اكتب اسمها في خانة عنوان فاضية")
areas = {
    "ليوا": ["الثروانية", "الجبانة", "مزيرعة", "حميم", "بوعوانة", "اليارية", "موقب", "الخيس",
             "نشاش", "وذين", "اليبانة", "يرّة", "موصل", "قرمدة", "نفير", "الرايقة", "المارية",
             "المارية الغربية", "عتاب", "خنّور", "الحميضي", "الصبخة", "حفيف", "قطوف", "الهيلة",
             "عرادة", "ملقطة", "العيد", "أم القرين", "العاطر", "الهذي", "المشرب", "مهدر بن عصيان"],
    "حبشان": ["حبشان الشرقية", "حبشان الغربية"],
    "الرويس": ["الرويس"],
    "المرفأ": ["المرفأ", "الطريف"],
    "غياثي": ["غياثي"],
    "السلع": ["السلع"],
    "بدع زايد": ["بدع زايد"],
}
AREA_COLS = 10
SITE_ROWS = 50
for j in range(AREA_COLS):
    col = j + 1
    names = list(areas.keys())
    hdr = ws.cell(11, col, names[j] if j < len(names) else None)
    hdr.font = f(11, True, WHITE); hdr.fill = fill(GREEN2); hdr.alignment = CENTER; hdr.border = BORDER
    sites = areas[names[j]] if j < len(names) else []
    for i in range(SITE_ROWS):
        c = ws.cell(12 + i, col, sites[i] if i < len(sites) else None)
        c.font = f(10, False, "1A3FA8"); c.fill = fill(INPUT if i >= len(sites) else WHITE)
        c.alignment = RIGHT; c.border = BORDER
    ws.column_dimensions[chr(64 + col)].width = 16
ws.cell(10, 1, "الخانات الصفراء فاضية وجاهزة للإضافة").font = f(9, False, "777777")

def list_block(col, title, items, rows, width=26):
    L = col
    h = ws[f"{L}11"]; h.value = title
    h.font = f(11, True, WHITE); h.fill = fill(GREEN2); h.alignment = CENTER; h.border = BORDER
    for i in range(rows):
        c = ws[f"{L}{12+i}"]
        c.value = items[i] if i < len(items) else None
        c.font = f(10, False, "1A3FA8"); c.fill = fill(WHITE if i < len(items) else INPUT)
        c.alignment = RIGHT; c.border = BORDER
    ws.column_dimensions[L].width = width

TYPES = ["حرق مخلفات", "تحميل مياه جوفية في صهاريج", "تصريف / كب صرف صحي", "رمي نفايات عامة",
         "تسرب زيوت أو مواد كيميائية", "حفر آبار غير مصرح", "لا توجد مشاهدات"]
list_block("L", "أنواع المشاهدات", TYPES, 15, 28)
ACTIONS = ["تم تحرير مخالفة", "مشاهدة فقط", "لا يوجد"]
list_block("N", "الإجراء", ACTIONS, 3, 18)
MONTHS = ["يناير", "فبراير", "مارس", "أبريل", "مايو", "يونيو", "يوليو", "أغسطس", "سبتمبر",
          "أكتوبر", "نوفمبر", "ديسمبر"]
list_block("P", "الأشهر", MONTHS, 12, 12)
PERIODS = [("الربع الأول", 1, 3), ("الربع الثاني", 4, 6), ("النصف الأول", 1, 6),
           ("الربع الثالث", 7, 9), ("الربع الرابع", 10, 12), ("النصف الثاني", 7, 12),
           ("التقرير السنوي", 1, 12)]
for i, (k, a, b) in enumerate([("الفترة", "من شهر", "إلى شهر")] + PERIODS):
    for j, v in enumerate((k, a, b)):
        c = ws.cell(11 + i, 18 + j, v)
        c.border = BORDER; c.alignment = CENTER
        c.font = f(11, True, WHITE) if i == 0 else f(10)
        if i == 0: c.fill = fill(GREEN2)
ws.column_dimensions["R"].width = 16; ws.column_dimensions["S"].width = 9; ws.column_dimensions["T"].width = 9

SCHED = [("اليوم", "المناطق المعتادة"), ("الاثنين", "ليوا (محضر أو محضرين) / المرفأ / الرويس / غياثي / السلع"),
         ("الثلاثاء", "حبشان (شرق أو غرب)"), ("الأربعاء", "ليوا (محضر أو محضرين) / المرفأ / الرويس / غياثي / السلع"),
         ("الخميس", "حبشان (شرق أو غرب)"), ("الجمعة", "بدع زايد (نصف دوام)")]
for i, (a, b) in enumerate(SCHED):
    ca = ws.cell(11 + i, 22, a); cb = ws.cell(11 + i, 23, b)
    for c in (ca, cb):
        c.border = BORDER; c.alignment = RIGHT
        c.font = f(11, True, WHITE) if i == 0 else f(10)
        if i == 0: c.fill = fill(GREEN2)
ws.column_dimensions["V"].width = 11; ws.column_dimensions["W"].width = 50
ws.column_dimensions["K"].width = 3; ws.column_dimensions["M"].width = 3
ws.column_dimensions["O"].width = 3; ws.column_dimensions["Q"].width = 3; ws.column_dimensions["U"].width = 3
ws.freeze_panes = "A12"

def name(n, ref):
    wb.defined_names[n] = DefinedName(n, attr_text=ref)

name("InspectorName", f"{S}!$B$3")
name("OrgName", f"{S}!$B$4")
name("SectorName", f"{S}!$B$5")
name("PhotoDir", f"{S}!$B$6")
name("AreaHdr", f"{S}!$A$11:$J$11")
name("SitesTop", f"{S}!$A$12")
name("AreaList", f"OFFSET({S}!$A$11,0,0,1,COUNTA({S}!$A$11:$J$11))")
name("TypeList", f"OFFSET({S}!$L$12,0,0,COUNTA({S}!$L$12:$L$26),1)")
name("ActionList", f"{S}!$N$12:$N$14")
name("MonthList", f"{S}!$P$12:$P$23")
name("PeriodList", f"{S}!$R$12:$R$18")
name("PeriodTbl", f"{S}!$R$12:$T$18")

# ------------------------------------------------------------------ log
ws = ws_log
cols = [  # (header, width, kind)  kind: in = input, fx = formula, hid = helper
    ("م", 5, "fx"), ("التاريخ", 12, "in"), ("المنطقة", 12, "in"), ("الموقع / المحضر", 18, "in"),
    ("وقت الوصول", 10, "in"), ("وقت المغادرة", 10, "in"), ("نوع المشاهدة", 24, "in"),
    ("وصف المشاهدة", 44, "in"), ("الإجراء", 16, "in"), ("عدد الصور", 9, "in"),
    ("خط العرض", 11, "in"), ("خط الطول", 11, "in"), ("ملاحظات", 26, "in"),
    ("اليوم", 10, "fx"), ("الخريطة", 12, "fx"), ("الصور", 14, "fx"),
    ("الشهر", 7, "hid"), ("السنة", 7, "hid"), ("بداية الأسبوع", 11, "hid"), ("زيارة جديدة", 7, "hid"),
    ("منطقة جديدة باليوم", 7, "hid"), ("مشاهدة", 7, "hid"), ("مخالفة", 7, "hid"), ("ترتيب الأسبوع", 7, "hid"),
]
banner(ws, "A1:P1", "سجل الزيارات التفتيشية")
ws["A2"] = "الأعمدة الخضراء تُعبّأ (يدويًا أو بلصق ما يصدّره التطبيق ابتداءً من عمود التاريخ). الأعمدة الرمادية تُحسب تلقائيًا ولا تحتاج تعديل."
ws.merge_cells("A2:P2"); ws["A2"].font = f(10, False, "555555"); ws["A2"].alignment = RIGHT
ws.merge_cells("Q3:X3"); ws["Q3"] = "أعمدة مساعدة — لا تعدّلها"
ws["Q3"].font = f(9, True, "777777"); ws["Q3"].alignment = CENTER
for j, (h, w, kind) in enumerate(cols):
    c = ws.cell(4, j + 1, h)
    c.font = f(11, True, WHITE if kind == "in" else "333333")
    c.fill = fill(GREEN2 if kind == "in" else ("B0BEC5" if kind == "fx" else "CFD8DC"))
    c.alignment = CENTER; c.border = BORDER
    ws.column_dimensions[c.column_letter].width = w
ws.row_dimensions[4].height = 32

sample = [
    (dt.date(2026, 10, 26), "ليوا", "الثروانية", dt.time(8, 10), dt.time(9, 25), "حرق مخلفات",
     "حرق مخلفات خشبية وسعف نخيل داخل مزرعة، الدخان واضح من الطريق", "تم تحرير مخالفة", 4, 23.1402, 53.7311, "صف تجريبي — احذفه قبل البدء"),
    (dt.date(2026, 10, 26), "ليوا", "الجبانة", dt.time(10, 5), dt.time(10, 50), "تحميل مياه جوفية في صهاريج",
     "صهريج يعبّي من بئر داخل مزرعة، لم يتواجد السائق", "مشاهدة فقط", 3, 23.1215, 53.6894, "صف تجريبي — احذفه قبل البدء"),
    (dt.date(2026, 10, 27), "حبشان", "حبشان الشرقية", dt.time(8, 30), dt.time(11, 0), "لا توجد مشاهدات",
     "جولة على الموقع بدون ملاحظات بيئية", "لا يوجد", 1, 23.8312, 53.6427, "صف تجريبي — احذفه قبل البدء"),
    (dt.date(2026, 10, 28), "المرفأ", "المرفأ", dt.time(8, 15), dt.time(9, 40), "تصريف / كب صرف صحي",
     "آثار تصريف صرف صحي من صهريج على أرض فضاء خلف المنطقة الصناعية", "مشاهدة فقط", 5, 24.0921, 53.4566, "صف تجريبي — احذفه قبل البدء"),
    (dt.date(2026, 10, 28), "الرويس", "الرويس", dt.time(11, 0), dt.time(12, 10), "رمي نفايات عامة",
     "رمي نفايات منزلية بجانب الطريق، تم ضبط المخالف", "تم تحرير مخالفة", 2, 24.1103, 52.7302, "صف تجريبي — احذفه قبل البدء"),
]
DATE_FMT = "yyyy/mm/dd"
if CLEAN: sample = []
for i in range(N):
    r = R0 + i
    if i < len(sample):
        for j, v in enumerate(sample[i]):
            ws.cell(r, 2 + j, v)
    ws.cell(r, 1, f'=IF($B{r}="","",ROW()-{R0-1})')
    ws.cell(r, 14, f'=IF($B{r}="","",CHOOSE(WEEKDAY($B{r}),"الأحد","الاثنين","الثلاثاء","الأربعاء","الخميس","الجمعة","السبت"))')
    ws.cell(r, 15, f'=IF(OR($K{r}="",$L{r}=""),"",HYPERLINK("https://maps.google.com/?q="&$K{r}&","&$L{r},"عرض الموقع"))')
    ws.cell(r, 16, f'=IF(OR($B{r}="",N($J{r})=0),"",HYPERLINK(PhotoDir&"/"&YEAR($B{r})&"-"&TEXT(MONTH($B{r}),"00")&"-"&TEXT(DAY($B{r}),"00")&"_"&$D{r},"فتح الصور ("&$J{r}&")"))')
    ws.cell(r, 17, f'=IF($B{r}="","",MONTH($B{r}))')
    ws.cell(r, 18, f'=IF($B{r}="","",YEAR($B{r}))')
    ws.cell(r, 19, f'=IF($B{r}="","",$B{r}-WEEKDAY($B{r},3))')
    ws.cell(r, 20, f'=IF($B{r}="",0,IF(COUNTIFS($B${R0}:$B{r},$B{r},$D${R0}:$D{r},$D{r})=1,1,0))')
    ws.cell(r, 21, f'=IF($B{r}="",0,IF(COUNTIFS($B${R0}:$B{r},$B{r},$C${R0}:$C{r},$C{r})=1,1,0))')
    ws.cell(r, 22, f'=IF(OR($B{r}="",$G{r}="",$G{r}="لا توجد مشاهدات"),0,1)')
    ws.cell(r, 23, f'=IF(AND($B{r}<>"",$I{r}="تم تحرير مخالفة"),1,0)')
    ws.cell(r, 24, f'=IF($B{r}="","",IF($S{r}=WkStart,COUNTIF($S${R0}:$S{r},WkStart),""))')
style(ws, f"A{R0}:P{R1}", f(10), None, RIGHT)
style(ws, f"Q{R0}:X{R1}", f(9, False, "888888"), fill(GREY), CENTER)
for r in range(R0, R1 + 1):
    ws.cell(r, 2).number_format = DATE_FMT
    ws.cell(r, 5).number_format = "hh:mm"; ws.cell(r, 6).number_format = "hh:mm"
    ws.cell(r, 11).number_format = "0.000000"; ws.cell(r, 12).number_format = "0.000000"
    ws.cell(r, 19).number_format = DATE_FMT
    for cc in (1, 14, 15, 16):
        ws.cell(r, cc).fill = fill("F5F7F8")
    ws.cell(r, 15).font = f(10, False, "1565C0"); ws.cell(r, 16).font = f(10, False, "1565C0")
    ws.cell(r, 1).alignment = CENTER; ws.cell(r, 10).alignment = CENTER
for L in "QRSTUVWX":
    ws.column_dimensions[L].hidden = True
ws.freeze_panes = "E5"
ws.auto_filter.ref = f"A4:P{R1}"

dv_area = DataValidation(type="list", formula1="=AreaList", allow_blank=True, showErrorMessage=True,
                         errorStyle="warning", errorTitle="منطقة غير موجودة",
                         error="المنطقة مو موجودة في القائمة. أضفها في شيت الإعدادات.")
dv_site = DataValidation(type="list",
                         formula1=f"=OFFSET(SitesTop,0,MATCH($C{R0},AreaHdr,0)-1,MAX(1,COUNTA(OFFSET(SitesTop,0,MATCH($C{R0},AreaHdr,0)-1,50,1))),1)",
                         allow_blank=True, showErrorMessage=True, errorStyle="warning",
                         errorTitle="موقع غير موجود", error="الموقع مو موجود تحت هذي المنطقة. أضفه في شيت الإعدادات عشان يظهر بالقائمة.")
dv_type = DataValidation(type="list", formula1="=TypeList", allow_blank=True, showErrorMessage=True,
                         errorStyle="warning", error="النوع مو موجود. أضفه في شيت الإعدادات.")
dv_act = DataValidation(type="list", formula1="=ActionList", allow_blank=True)
dv_date = DataValidation(type="date", operator="greaterThan", formula1="DATE(2020,1,1)", allow_blank=True,
                         showErrorMessage=True, error="اكتب تاريخ صحيح")
for dv, col in ((dv_area, "C"), (dv_site, "D"), (dv_type, "G"), (dv_act, "I"), (dv_date, "B")):
    dv.add(f"{col}{R0}:{col}{R1}")
    ws.add_data_validation(dv)

ws.conditional_formatting.add(f"I{R0}:I{R1}", FormulaRule(formula=[f'$I{R0}="تم تحرير مخالفة"'], fill=fill("FDE2E1"), font=Font(name=FONT, color="B3261E", bold=True)))
ws.conditional_formatting.add(f"I{R0}:I{R1}", FormulaRule(formula=[f'$I{R0}="مشاهدة فقط"'], fill=fill("FFF1D6"), font=Font(name=FONT, color="8A5A00")))
ws.conditional_formatting.add(f"I{R0}:I{R1}", FormulaRule(formula=[f'$I{R0}="لا يوجد"'], fill=fill("E3F2E8"), font=Font(name=FONT, color="1E5B3C")))
ws.conditional_formatting.add(f"A{R0}:P{R1}", FormulaRule(formula=[f'$M{R0}="صف تجريبي — احذفه قبل البدء"'], fill=fill("F3E5F5")))

def lname(n, col):
    name(n, f"{LOG}!${col}${R0}:${col}${R1}")
for n, c in [("L_Date", "B"), ("L_Area", "C"), ("L_Site", "D"), ("L_Arr", "E"), ("L_Type", "G"), ("L_Desc", "H"),
             ("L_Act", "I"), ("L_Pics", "J"), ("L_Month", "Q"), ("L_Year", "R"), ("L_Week", "S"),
             ("L_NewVisit", "T"), ("L_NewArea", "U"), ("L_Obs", "V"), ("L_Viol", "W"), ("L_WIdx", "X")]:
    lname(n, c)

# ------------------------------------------------------------------ shared report helpers
WID = {"A": 2, "B": 12, "C": 12, "D": 20, "E": 34, "F": 15, "G": 14, "H": 14, "I": 11, "J": 2}
def setup_report(ws, title):
    for k, v in WID.items():
        ws.column_dimensions[k].width = v
    banner(ws, "B1:I1", title, 18)
    ws.row_dimensions[1].height = 34
    ws.merge_cells("B2:I2")
    ws["B2"] = '=OrgName&"  |  "&SectorName'
    style(ws, "B2:I2", f(11, True, GREEN), fill(SAND), CENTER, border=False)
    ws.page_setup.orientation = "portrait"; ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4

def input_cell(ws, ref, value, fmt=None):
    c = ws[ref]; c.value = value
    c.font = f(12, True, "1A3FA8"); c.fill = fill(INPUT); c.alignment = CENTER; c.border = BORDER
    if fmt: c.number_format = fmt

def label(ws, ref, text, merge=None):
    if merge: ws.merge_cells(merge)
    c = ws[ref]; c.value = text; c.font = f(11, True, "333333"); c.alignment = RIGHT

def kpis(ws, row, items):
    """items: list of (merge range cols e.g. ('B','C'), label, formula, color)"""
    for (a, b), lab, fx, color in items:
        ws.merge_cells(f"{a}{row}:{b}{row}")
        ws.merge_cells(f"{a}{row+1}:{b}{row+1}")
        ws[f"{a}{row}"] = lab
        ws[f"{a}{row+1}"] = fx
        style(ws, f"{a}{row}:{b}{row}", f(10, True, WHITE), fill(color), CENTER)
        style(ws, f"{a}{row+1}:{b}{row+1}", f(22, True, color), fill(WHITE), CENTER)
    ws.row_dimensions[row].height = 22
    ws.row_dimensions[row + 1].height = 42

def header_row(ws, row, headers):
    for col, h in headers:
        c = ws[f"{col}{row}"]; c.value = h
        c.font = f(10, True, WHITE); c.fill = fill(GREEN2); c.alignment = CENTER; c.border = BORDER
    ws.row_dimensions[row].height = 26

def signature(ws, row):
    ws.merge_cells(f"B{row}:E{row}"); ws.merge_cells(f"F{row}:I{row}")
    ws[f"B{row}"] = '="المفتش: "&InspectorName'
    ws[f"F{row}"] = "التوقيع: ............................"
    ws.merge_cells(f"B{row+1}:E{row+1}")
    ws[f"B{row+1}"] = '="تاريخ الإعداد: "&TEXT(TODAY(),"yyyy/mm/dd")'
    for ref in (f"B{row}", f"F{row}", f"B{row+1}"):
        ws[ref].font = f(11, True, "333333"); ws[ref].alignment = RIGHT
    ws.row_dimensions[row].height = 30

C_V, C_O, C_X, C_A, C_OK = "1E5B3C", "C77700", "B3261E", "8A5A00", "2E6E8E"

# ------------------------------------------------------------------ weekly
ws = ws_w
setup_report(ws, "التقرير الأسبوعي للتفتيش البيئي")
label(ws, "B4", "اكتب أي تاريخ من الأسبوع:", "B4:C4")
input_cell(ws, "D4", dt.date(2026, 10, 28), DATE_FMT)
ws["E4"] = "← يحسب الملف بداية ونهاية الأسبوع تلقائيًا"
ws["E4"].font = f(9, False, "777777"); ws["E4"].alignment = RIGHT
label(ws, "B5", "الفترة:", "B5:C5")
ws["D5"] = "=D4-WEEKDAY(D4,3)"; ws["D5"].number_format = DATE_FMT
ws["E5"] = '="من الاثنين "&TEXT(D5,"yyyy/mm/dd")&"  إلى الجمعة "&TEXT(D5+4,"yyyy/mm/dd")'
ws["F5"] = '="الأسبوع رقم "&WEEKNUM(D5,2)'
label(ws, "B6", "المفتش:", "B6:C6"); ws["D6"] = "=InspectorName"
for ref in ("D5", "E5", "F5", "D6"):
    ws[ref].font = f(11, True, GREEN); ws[ref].alignment = RIGHT
name("WkStart", "'التقرير الأسبوعي'!$D$5")

kpis(ws, 8, [
    (("B", "C"), "الزيارات", "=SUMIFS(L_NewVisit,L_Week,WkStart)", C_V),
    (("D", "D"), "المشاهدات البيئية", "=SUMIFS(L_Obs,L_Week,WkStart)", C_O),
    (("E", "E"), "المخالفات المحررة", "=SUMIFS(L_Viol,L_Week,WkStart)", C_X),
    (("F", "G"), "مشاهدات دون مخالف", '=COUNTIFS(L_Week,WkStart,L_Act,"مشاهدة فقط")', C_A),
    (("H", "I"), "زيارات بلا مشاهدات", '=COUNTIFS(L_Week,WkStart,L_Type,"لا توجد مشاهدات")', C_OK),
])

section(ws, "B11:I11", "ملخص أيام الأسبوع")
header_row(ws, 12, [("B", "اليوم"), ("C", "التاريخ"), ("D", "المناطق"), ("E", "المواقع المزارة"),
                    ("F", "الزيارات"), ("G", "المشاهدات"), ("H", "المخالفات"), ("I", "دون مخالف")])
days = ["الاثنين", "الثلاثاء", "الأربعاء", "الخميس", "الجمعة"]
for i, d in enumerate(days):
    r = 13 + i
    ws[f"B{r}"] = d
    ws[f"C{r}"] = f"=WkStart+{i}"; ws[f"C{r}"].number_format = DATE_FMT
    ws[f"D{r}"] = ArrayFormula(f"D{r}", f'=_xlfn.TEXTJOIN("، ",TRUE,IF((L_Date=C{r})*(L_NewArea=1),L_Area,""))')
    ws[f"E{r}"] = ArrayFormula(f"E{r}", f'=_xlfn.TEXTJOIN("، ",TRUE,IF((L_Date=C{r})*(L_NewVisit=1),L_Site,""))')
    ws[f"F{r}"] = f"=SUMIFS(L_NewVisit,L_Date,C{r})"
    ws[f"G{r}"] = f"=SUMIFS(L_Obs,L_Date,C{r})"
    ws[f"H{r}"] = f"=SUMIFS(L_Viol,L_Date,C{r})"
    ws[f"I{r}"] = f'=COUNTIFS(L_Date,C{r},L_Act,"مشاهدة فقط")'
    style(ws, f"B{r}:I{r}", f(10), fill(WHITE if i % 2 == 0 else "F6FAF7"), CENTER)
    ws[f"C{r}"].number_format = DATE_FMT
    ws.row_dimensions[r].height = 30
r = 18
ws[f"B{r}"] = "الإجمالي"; ws.merge_cells(f"B{r}:E{r}")
for L in "FGHI":
    ws[f"{L}{r}"] = f"=SUM({L}13:{L}17)"
style(ws, f"B{r}:I{r}", f(11, True, GREEN), fill(LIGHT), CENTER)

def type_table(ws, top, crit):
    """crit: extra COUNTIFS criteria string e.g. 'L_Week,WkStart'"""
    header_row(ws, top, [("B", "نوع المشاهدة"), ("F", "العدد"), ("G", "مخالفة محررة"), ("H", "دون مخالف"), ("I", "النسبة")])
    ws.merge_cells(f"B{top}:E{top}")
    first = top + 1
    for i in range(10):
        r = first + i
        ws.merge_cells(f"B{r}:E{r}")
        ws[f"B{r}"] = f"=IF(INDEX({S}!$L$12:$L$26,{i+1})=\"\",\"\",INDEX({S}!$L$12:$L$26,{i+1}))"
        ws[f"F{r}"] = f'=IF($B{r}="","",COUNTIFS({crit},L_Type,$B{r}))'
        ws[f"G{r}"] = f'=IF($B{r}="","",COUNTIFS({crit},L_Type,$B{r},L_Act,"تم تحرير مخالفة"))'
        ws[f"H{r}"] = f'=IF($B{r}="","",COUNTIFS({crit},L_Type,$B{r},L_Act,"مشاهدة فقط"))'
        ws[f"I{r}"] = f'=IF(OR($B{r}="",SUM($F${first}:$F${first+9})=0),"",F{r}/SUM($F${first}:$F${first+9}))'
        style(ws, f"B{r}:I{r}", f(10), fill(WHITE if i % 2 == 0 else "F6FAF7"), CENTER)
        ws[f"B{r}"].alignment = RIGHT
        ws[f"I{r}"].number_format = "0%"
    return first, first + 9

type_first, type_last = None, None
section(ws, "B20:I20", "المشاهدات حسب النوع")
type_table(ws, 21, "L_Week,WkStart")

section(ws, "B33:I33", "تفاصيل زيارات الأسبوع")
header_row(ws, 34, [("B", "م"), ("C", "التاريخ"), ("D", "الموقع"), ("E", "الوصف"),
                    ("F", "نوع المشاهدة"), ("G", "الإجراء"), ("H", "وقت الوصول"), ("I", "الصور")])
DET = 40
for i in range(DET):
    r = 35 + i
    k = i + 1
    m = f"MATCH({k},L_WIdx,0)"
    ws[f"B{r}"] = f'=IF(ISNUMBER({m}),{k},"")'
    ws[f"C{r}"] = f'=IF($B{r}="","",INDEX(L_Date,{m}))'
    ws[f"D{r}"] = f'=IF($B{r}="","",INDEX(L_Area,{m})&" - "&INDEX(L_Site,{m}))'
    ws[f"E{r}"] = f'=IF($B{r}="","",INDEX(L_Desc,{m})&"")'
    ws[f"F{r}"] = f'=IF($B{r}="","",INDEX(L_Type,{m})&"")'
    ws[f"G{r}"] = f'=IF($B{r}="","",INDEX(L_Act,{m})&"")'
    ws[f"H{r}"] = f'=IF($B{r}="","",INDEX(L_Arr,{m}))'
    ws[f"I{r}"] = f'=IF($B{r}="","",N(INDEX(L_Pics,{m})))'
    style(ws, f"B{r}:I{r}", f(10), fill(WHITE if i % 2 == 0 else "F6FAF7"), CENTER)
    ws[f"E{r}"].alignment = RIGHT
    ws[f"C{r}"].number_format = DATE_FMT; ws[f"H{r}"].number_format = "hh:mm"
ws.conditional_formatting.add("G35:G74", FormulaRule(formula=['$G35="تم تحرير مخالفة"'], font=Font(name=FONT, color="B3261E", bold=True)))
ws["B75"] = "يعرض الجدول أول 40 سجل في الأسبوع."
ws["B75"].font = f(9, False, "888888")
section(ws, "B76:I76", "ملاحظات وتوصيات المفتش")
ws.merge_cells("B77:I80")
style(ws, "B77:I80", f(11, False, "1A3FA8"), fill(INPUT), Alignment(horizontal="right", vertical="top", wrap_text=True))
signature(ws, 82)
ws.print_area = "B1:I83"

# ------------------------------------------------------------------ monthly
ws = ws_m
setup_report(ws, "التقرير الشهري للتفتيش البيئي")
label(ws, "B4", "السنة:", "B4:C4"); input_cell(ws, "D4", 2026)
label(ws, "E4", "الشهر:"); input_cell(ws, "F4", "أكتوبر")
ws["E4"].alignment = Alignment(horizontal="left", vertical="center")
dv = DataValidation(type="list", formula1="=MonthList", allow_blank=False); dv.add("F4"); ws.add_data_validation(dv)
ws["H4"] = "=MATCH(F4,MonthList,0)"; ws["H4"].font = f(9, False, "BBBBBB")
ws["I4"] = '=DATE(D4,H4,1)'; ws["I4"].number_format = DATE_FMT; ws["I4"].font = f(9, False, "BBBBBB")
label(ws, "B5", "الفترة:", "B5:C5")
ws.merge_cells("D5:G5")
ws["D5"] = '="من "&TEXT(I4,"yyyy/mm/dd")&"  إلى "&TEXT(EOMONTH(I4,0),"yyyy/mm/dd")'
label(ws, "B6", "المفتش:", "B6:C6"); ws["D6"] = "=InspectorName"
for ref in ("D5", "D6"):
    ws[ref].font = f(11, True, GREEN); ws[ref].alignment = RIGHT
MC = "L_Year,$D$4,L_Month,$H$4"
kpis(ws, 8, [
    (("B", "C"), "الزيارات", f"=SUMIFS(L_NewVisit,{MC})", C_V),
    (("D", "D"), "المشاهدات البيئية", f"=SUMIFS(L_Obs,{MC})", C_O),
    (("E", "E"), "المخالفات المحررة", f"=SUMIFS(L_Viol,{MC})", C_X),
    (("F", "G"), "مشاهدات دون مخالف", f'=COUNTIFS({MC},L_Act,"مشاهدة فقط")', C_A),
    (("H", "I"), "زيارات بلا مشاهدات", f'=COUNTIFS({MC},L_Type,"لا توجد مشاهدات")', C_OK),
])

def area_table(ws, top, sum_crit):
    header_row(ws, top, [("B", "المنطقة"), ("D", "الزيارات"), ("E", "المشاهدات البيئية"),
                         ("F", "مخالفات محررة"), ("G", "دون مخالف"), ("H", "بلا مشاهدات"), ("I", "النسبة")])
    ws.merge_cells(f"B{top}:C{top}")
    first = top + 1
    for i in range(AREA_COLS):
        r = first + i
        ws.merge_cells(f"B{r}:C{r}")
        ws[f"B{r}"] = f'=IF(INDEX(AreaHdr,{i+1})="","",INDEX(AreaHdr,{i+1}))'
        ws[f"D{r}"] = f'=IF($B{r}="","",SUMIFS(L_NewVisit,{sum_crit},L_Area,$B{r}))'
        ws[f"E{r}"] = f'=IF($B{r}="","",SUMIFS(L_Obs,{sum_crit},L_Area,$B{r}))'
        ws[f"F{r}"] = f'=IF($B{r}="","",SUMIFS(L_Viol,{sum_crit},L_Area,$B{r}))'
        ws[f"G{r}"] = f'=IF($B{r}="","",COUNTIFS({sum_crit},L_Area,$B{r},L_Act,"مشاهدة فقط"))'
        ws[f"H{r}"] = f'=IF($B{r}="","",COUNTIFS({sum_crit},L_Area,$B{r},L_Type,"لا توجد مشاهدات"))'
        ws[f"I{r}"] = f'=IF(OR($B{r}="",SUM($D${first}:$D${first+AREA_COLS-1})=0),"",D{r}/SUM($D${first}:$D${first+AREA_COLS-1}))'
        style(ws, f"B{r}:I{r}", f(10), fill(WHITE if i % 2 == 0 else "F6FAF7"), CENTER)
        ws[f"I{r}"].number_format = "0%"
    r = first + AREA_COLS
    ws.merge_cells(f"B{r}:C{r}"); ws[f"B{r}"] = "الإجمالي"
    for L in "DEFGH":
        ws[f"{L}{r}"] = f"=SUM({L}{first}:{L}{r-1})"
    style(ws, f"B{r}:I{r}", f(11, True, GREEN), fill(LIGHT), CENTER)
    return r

section(ws, "B11:I11", "حسب المنطقة")
end = area_table(ws, 12, MC)            # rows 12..23
section(ws, "B25:I25", "حسب الأسبوع")
header_row(ws, 26, [("B", "الأسبوع"), ("E", "الفترة"), ("F", "الزيارات"), ("G", "المشاهدات"), ("H", "المخالفات"), ("I", "دون مخالف")])
ws.merge_cells("B26:D26")
for i in range(6):
    r = 27 + i
    ws.merge_cells(f"B{r}:D{r}")
    mon = f"($I$4-WEEKDAY($I$4,3)+{7*i})"
    ws[f"B{r}"] = f'=IF({mon}>EOMONTH($I$4,0),"","الأسبوع "&{i+1})'
    ws[f"E{r}"] = f'=IF($B{r}="","",TEXT(MAX({mon},$I$4),"mm/dd")&" - "&TEXT(MIN({mon}+6,EOMONTH($I$4,0)),"mm/dd"))'
    ws[f"F{r}"] = f'=IF($B{r}="","",SUMIFS(L_NewVisit,L_Week,{mon},{MC}))'
    ws[f"G{r}"] = f'=IF($B{r}="","",SUMIFS(L_Obs,L_Week,{mon},{MC}))'
    ws[f"H{r}"] = f'=IF($B{r}="","",SUMIFS(L_Viol,L_Week,{mon},{MC}))'
    ws[f"I{r}"] = f'=IF($B{r}="","",COUNTIFS(L_Week,{mon},{MC},L_Act,"مشاهدة فقط"))'
    style(ws, f"B{r}:I{r}", f(10), fill(WHITE if i % 2 == 0 else "F6FAF7"), CENTER)
section(ws, "B34:I34", "حسب نوع المشاهدة")
tf, tl = type_table(ws, 35, MC)          # rows 35..45

def type_chart(ws, tf, tl, anchor, title):
    ch = BarChart(); ch.type = "bar"; ch.style = 10
    ch.title = title; ch.legend = None
    data = Reference(ws, min_col=6, min_row=tf - 1, max_row=tl)
    cats = Reference(ws, min_col=2, min_row=tf, max_row=tl)
    ch.add_data(data, titles_from_data=True); ch.set_categories(cats)
    ch.series[0].graphicalProperties.solidFill = C_O
    ch.height = 7.5; ch.width = 17
    ws.add_chart(ch, anchor)

type_chart(ws, tf, tl, "B47", "المشاهدات حسب النوع")
signature(ws, 65)
ws.print_area = "B1:I66"

# ------------------------------------------------------------------ periodic
ws = ws_p
setup_report(ws, "التقرير الدوري للتفتيش البيئي")
label(ws, "B4", "السنة:", "B4:C4"); input_cell(ws, "D4", 2026)
label(ws, "E4", "الفترة:"); input_cell(ws, "F4", "الربع الرابع")
ws["E4"].alignment = Alignment(horizontal="left", vertical="center")
dv = DataValidation(type="list", formula1="=PeriodList", allow_blank=False); dv.add("F4"); ws.add_data_validation(dv)
ws["H4"] = "=VLOOKUP(F4,PeriodTbl,2,FALSE)"; ws["I4"] = "=VLOOKUP(F4,PeriodTbl,3,FALSE)"
ws["H4"].font = ws["I4"].font = f(9, False, "BBBBBB")
ws["H4"].font = f(9, False, "BBBBBB")
ws["B1"] = '="تقرير "&F4&" للتفتيش البيئي — "&D4'
label(ws, "B5", "الفترة:", "B5:C5"); ws.merge_cells("D5:G5")
ws["D5"] = '="من "&TEXT(DATE(D4,H4,1),"yyyy/mm/dd")&"  إلى "&TEXT(EOMONTH(DATE(D4,I4,1),0),"yyyy/mm/dd")'
label(ws, "B6", "المفتش:", "B6:C6"); ws["D6"] = "=InspectorName"
for ref in ("D5", "D6"):
    ws[ref].font = f(11, True, GREEN); ws[ref].alignment = RIGHT
PC = 'L_Year,$D$4,L_Month,">="&$H$4,L_Month,"<="&$I$4'
kpis(ws, 8, [
    (("B", "C"), "الزيارات", f"=SUMIFS(L_NewVisit,{PC})", C_V),
    (("D", "D"), "المشاهدات البيئية", f"=SUMIFS(L_Obs,{PC})", C_O),
    (("E", "E"), "المخالفات المحررة", f"=SUMIFS(L_Viol,{PC})", C_X),
    (("F", "G"), "مشاهدات دون مخالف", f'=COUNTIFS({PC},L_Act,"مشاهدة فقط")', C_A),
    (("H", "I"), "زيارات بلا مشاهدات", f'=COUNTIFS({PC},L_Type,"لا توجد مشاهدات")', C_OK),
])
section(ws, "B11:I11", "حسب الشهر")
header_row(ws, 12, [("B", "الشهر"), ("D", "الزيارات"), ("E", "المشاهدات البيئية"), ("F", "مخالفات محررة"),
                    ("G", "دون مخالف"), ("H", "بلا مشاهدات"), ("I", "المواقع")])
ws.merge_cells("B12:C12")
for i in range(12):
    r = 13 + i
    mth = f"($H$4+{i})"
    ws.merge_cells(f"B{r}:C{r}")
    ws[f"B{r}"] = f'=IF({mth}>$I$4,"",INDEX(MonthList,{mth}))'
    crit = f"L_Year,$D$4,L_Month,{mth}"
    ws[f"D{r}"] = f'=IF($B{r}="","",SUMIFS(L_NewVisit,{crit}))'
    ws[f"E{r}"] = f'=IF($B{r}="","",SUMIFS(L_Obs,{crit}))'
    ws[f"F{r}"] = f'=IF($B{r}="","",SUMIFS(L_Viol,{crit}))'
    ws[f"G{r}"] = f'=IF($B{r}="","",COUNTIFS({crit},L_Act,"مشاهدة فقط"))'
    ws[f"H{r}"] = f'=IF($B{r}="","",COUNTIFS({crit},L_Type,"لا توجد مشاهدات"))'
    ws[f"I{r}"] = f'=IF($B{r}="","",SUMIFS(L_Pics,{crit}))'
    style(ws, f"B{r}:I{r}", f(10), fill(WHITE if i % 2 == 0 else "F6FAF7"), CENTER)
ws["I12"] = "الصور"
r = 25
ws.merge_cells(f"B{r}:C{r}"); ws[f"B{r}"] = "الإجمالي"
for L in "DEFGHI":
    ws[f"{L}{r}"] = f"=SUM({L}13:{L}24)"
style(ws, f"B{r}:I{r}", f(11, True, GREEN), fill(LIGHT), CENTER)

ch = BarChart(); ch.type = "col"; ch.grouping = "clustered"; ch.style = 10
ch.title = "المشاهدات والمخالفات حسب الشهر"
data = Reference(ws, min_col=5, max_col=6, min_row=12, max_row=24)
cats = Reference(ws, min_col=2, min_row=13, max_row=24)
ch.add_data(data, titles_from_data=True); ch.set_categories(cats)
ch.series[0].graphicalProperties.solidFill = C_O
ch.series[1].graphicalProperties.solidFill = C_X
ch.height = 7.5; ch.width = 17
ws.add_chart(ch, "B27")

section(ws, "B43:I43", "حسب المنطقة")
area_table(ws, 44, PC)                 # 44..55
section(ws, "B57:I57", "حسب نوع المشاهدة")
tf, tl = type_table(ws, 58, PC)        # 58..68
type_chart(ws, tf, tl, "B70", "المشاهدات حسب النوع")
section(ws, "B86:I86", "ملاحظات وتوصيات المفتش")
ws.merge_cells("B87:I91")
ws["B87"].fill = fill(INPUT); ws["B87"].alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
ws["B87"].font = f(11, False, "1A3FA8")
style(ws, "B87:I91", None, fill(INPUT), None)
signature(ws, 93)
ws.print_area = "B1:I94"

# notes area on monthly too
section(ws_m, "B62:I62", "ملاحظات المفتش")
ws_m.merge_cells("B63:I63"); style(ws_m, "B63:I63", f(11, False, "1A3FA8"), fill(INPUT), Alignment(horizontal="right", vertical="top", wrap_text=True))
ws_m.row_dimensions[63].height = 60

# ------------------------------------------------------------------ help
ws = ws_help
ws.column_dimensions["A"].width = 3; ws.column_dimensions["B"].width = 5; ws.column_dimensions["C"].width = 95
banner(ws, "B1:C1", "سجل التفتيش البيئي — طريقة الاستخدام", 18)
ws.row_dimensions[1].height = 36
steps = [
    ("", "فكرة الملف", True),
    ("١", "كل صف في «سجل الزيارات» = مشاهدة واحدة في موقع واحد. لو في الموقع أكثر من مشاهدة، سجّل كل وحدة في صف بنفس التاريخ والموقع، والملف يحسبها زيارة واحدة تلقائيًا.", False),
    ("٢", "لو زرت موقع وما لقيت شيء، سجّل صف واحد ونوع المشاهدة «لا توجد مشاهدات» والإجراء «لا يوجد».", False),
    ("٣", "التقارير الأسبوعية والشهرية والدورية تُحسب كلها تلقائيًا من السجل — ما تكتب فيها شيء غير الخانات الصفراء (التاريخ أو الشهر أو الفترة).", False),
    ("", "الإدخال", True),
    ("٤", "عبّئ الأعمدة الخضراء: التاريخ، المنطقة، الموقع، الأوقات، نوع المشاهدة، الوصف، الإجراء، عدد الصور، الإحداثيات، الملاحظات.", False),
    ("٥", "المنطقة والموقع ونوع المشاهدة والإجراء كلها قوائم تختار منها. قائمة المواقع تتغير حسب المنطقة اللي اخترتها.", False),
    ("٦", "من التطبيق: اضغط «تصدير»، افتح ملف الإكسل اللي يطلعه، انسخ الصفوف، والصقها في أول صف فاضي في السجل ابتداءً من عمود «التاريخ» (لصق قيم فقط).", False),
    ("", "الصور", True),
    ("٧", "الصور ما تنحط داخل الإكسل عشان الملف ما يثقل. تكون في فولدر اسمه «الصور» بجانب ملف الإكسل، وداخله فولدر لكل زيارة باسم: التاريخ_الموقع (مثال: 2026-10-26_الثروانية).", False),
    ("٨", "عمود «الصور» في السجل فيه رابط يفتح فولدر الزيارة مباشرة، وعمود «الخريطة» يفتح موقع الزيارة على خرائط جوجل.", False),
    ("", "الإعدادات", True),
    ("٩", "في شيت «الإعدادات» تقدر تغيّر اسم المفتش والجهة، وتضيف مواقع أو مناطق أو أنواع مشاهدات جديدة في الخانات الصفراء — وتظهر فورًا في القوائم والتقارير.", False),
    ("", "قبل البدء", True),
    ("١٠", "في السجل ٥ صفوف تجريبية ملوّنة بالبنفسجي عشان تشوف شكل التقارير. احذف محتواها (من عمود التاريخ إلى الملاحظات) قبل ما تبدأ التسجيل الفعلي.", False),
]
r = 3
for num, text, is_head in steps:
    if is_head:
        ws.merge_cells(f"B{r}:C{r}"); ws[f"B{r}"] = text
        style(ws, f"B{r}:C{r}", f(13, True, GREEN), fill(LIGHT), RIGHT, border=False)
        ws.row_dimensions[r].height = 26
    else:
        ws[f"B{r}"] = num; ws[f"C{r}"] = text
        ws[f"B{r}"].font = f(12, True, GREEN2); ws[f"B{r}"].alignment = CENTER
        ws[f"C{r}"].font = f(11); ws[f"C{r}"].alignment = RIGHT
        ws.row_dimensions[r].height = 36
    r += 1
ws[f"B{r+1}"] = "مفتاح الألوان:"; ws[f"B{r+1}"].font = f(11, True)
ws.merge_cells(f"B{r+1}:C{r+1}")
legend = [("FFF9DB", "خانة تكتب فيها أنت"), (GREEN2, "عمود إدخال في السجل"), ("B0BEC5", "عمود يُحسب تلقائيًا")]
for i, (c, t) in enumerate(legend):
    rr = r + 2 + i
    ws[f"B{rr}"].fill = fill(c); ws[f"B{rr}"].border = BORDER
    ws[f"C{rr}"] = t; ws[f"C{rr}"].font = f(10); ws[f"C{rr}"].alignment = RIGHT

wb.active = 2 if CLEAN else 0
if CLEAN:
    for w in wb.worksheets: w.sheet_view.tabSelected = (w.title == 'التقرير الأسبوعي')
wb.save(OUT)
print("saved", OUT)
